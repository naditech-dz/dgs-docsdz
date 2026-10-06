#!/usr/bin/env python3
"""Génère src/data.js (catalogue des champs + liste des documents) à partir de
fields_catalog.json, INDEX-DGS-DOCUMENTS.xlsx et du dossier des modèles."""
import json, os, re, sys, glob
from openpyxl import load_workbook

here = os.path.dirname(os.path.abspath(__file__))
tpl_root = sys.argv[1]  # dossier DGS_DOCUMENTS_V1.0
cat = json.load(open(os.path.join(here, 'fields_catalog.json'), encoding='utf8'))

# titres propres depuis l'index
titles = {}
wb = load_workbook(os.path.join(here, 'INDEX-DGS-DOCUMENTS.xlsx'), data_only=True)
for r in list(wb['INDEX'].iter_rows(values_only=True))[4:]:
    if r[0] and r[1]:
        titles[str(r[0]).strip()] = {'t': str(r[1]).strip(), 'u': (r[3] or ''), 's': (r[4] or '')}

docs = []
for f in sorted(glob.glob(os.path.join(tpl_root, '**/*.docx'), recursive=True)):
    rel = os.path.relpath(f, tpl_root).replace(os.sep, '/')
    if rel.startswith('00_README'):
        continue
    m = re.match(r'(DGS-[A-Z]+-\d+)', os.path.basename(rel))
    if not m:
        continue
    code = m.group(1)
    ti = titles.get(code, {})
    name = re.sub(r'^DGS-[A-Z]+-\d+-', '', os.path.splitext(os.path.basename(rel))[0]).replace('-', ' ')
    docs.append({'c': code, 'p': rel, 'f': os.path.basename(rel), 't': ti.get('t', name), 'u': ti.get('u', ''), 's': ti.get('s', '')})

fields = {}   # key -> label
docfields = {}  # code -> [[key, count], ...]
for code, items in cat['catalog'].items():
    lst = []
    for it in items:
        k = it['key']
        t = it['text'][1:-1].strip()
        if k not in fields:
            fields[k] = t[:1].upper() + t[1:].lower()
        lst.append([k, it['count']])
    docfields[code] = lst

out = 'const DOCS=%s;\nconst FIELDS=%s;\nconst DOCFIELDS=%s;\n' % (
    json.dumps(docs, ensure_ascii=False, separators=(',', ':')),
    json.dumps(fields, ensure_ascii=False, separators=(',', ':')),
    json.dumps(docfields, ensure_ascii=False, separators=(',', ':')))
open(os.path.join(here, '..', 'src', 'data.js'), 'w', encoding='utf8').write(out)
print(len(docs), 'documents,', len(fields), 'champs ->', len(out), 'octets')
